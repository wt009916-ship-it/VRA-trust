"""Durable project inbox: source-first extraction, explicit review, atomic adoption.

Parsing is deterministic and local. Text extraction is not engineering validation.
No external LLM receives documents through this workflow.
"""
import csv
import hashlib
import io
import json
import re
import shutil
import threading
from collections import Counter
from pathlib import Path

from fastapi import BackgroundTasks, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from pydantic import Field, StrictInt, field_validator

from . import core
from .domain import identifier
from .schema import Contract, Evidence, Project, ProjectCreate
from .weather import AnalysisOptions, analyze_epw

TYPES = {'.pdf': 'drawing', '.ifc': 'bim', '.dxf': 'drawing', '.dwg': 'drawing',
         '.png': 'drawing', '.jpg': 'drawing', '.jpeg': 'drawing', '.idf': 'idf',
         '.epw': 'epw', '.csv': 'equipment_table', '.xlsx': 'equipment_table',
         '.json': 'manual_input', '.txt': 'manual_input'}
FIELDS = {'name': '项目名称', 'use': '建筑用途', 'location': '项目地点',
          'region': '地区代码', 'area_m2': '总建筑面积'}


class InboxCreate(Contract):
    project_id: str | None = None


class InboxCommit(Contract):
    project: ProjectCreate | None = None
    reviewed: bool


class InboxDraft(Contract):
    expected_revision: StrictInt = Field(ge=0)
    fields: dict[str, str] = Field(max_length=5)

    @field_validator('fields')
    @classmethod
    def bounded_fields(cls, fields):
        limits = {'name': 160, 'use': 100, 'location': 200, 'region': 40, 'area_m2': 32}
        if any(k not in limits or len(v) > limits[k] for k, v in fields.items()):
            raise ValueError('草稿仅包含项目名称、用途、地点、地区和面积，且不能超过字段长度限制')
        return fields


class TextCorrection(Contract):
    expected_revision: StrictInt = Field(ge=0)
    page: StrictInt = Field(ge=1, le=80)
    text: str = Field(max_length=30000)
    reason: str = Field(min_length=3, max_length=500)


def text_candidates(text, source):
    """Only explicit labelled declarations. Never infer gross area from partial rooms."""
    found = []
    for field, label in FIELDS.items():
        pattern = rf'(?:^|\n)\s*(?:{label}|{field})\s*[:：=]\s*([^\n]+)'
        for match in re.finditer(pattern, text, re.I):
            value = match[1].strip()
            if field == 'area_m2':
                number = re.fullmatch(r'((?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(?:m2|m²|㎡|平方米)?', value, re.I)
                if not number:
                    continue
                value = float(number[1].replace(',', ''))
                if not 0 < value <= 10000000:
                    continue
            found.append({'field': field, 'value': value, 'locator': source,
                          'excerpt': match[0].strip()[:300]})
    return found


def parse_source(path):
    ext = path.suffix.lower()
    out = {'parser': 'local-intake-1', 'candidates': [], 'facts': {}, 'warnings': [], 'next_view': 'evidence'}
    if ext == '.epw':
        result = analyze_epw(path, AnalysisOptions())
        out.update(facts={k: result[k] for k in ('location', 'quality', 'annual')}, next_view='weather')
        out['warnings'] = result['quality']['issues'] + ['气象站地点不会自动填写为建筑地点。']
    elif ext == '.ifc':
        import ifcopenshell
        model = ifcopenshell.open(str(path))
        spaces = model.by_type('IfcSpace')
        out['facts'] = {'schema': model.schema, 'spaces': len(spaces), 'storeys': len(model.by_type('IfcBuildingStorey')),
                        'spaces_with_body': sum(any(r.RepresentationIdentifier == 'Body' for r in getattr(s.Representation, 'Representations', [])) for s in spaces)}
        out['next_view'] = 'spatial'
        out['warnings'] = ['实体清点不等于热模型完成；空间几何、用途和边界仍需识别及复核。']
    elif ext == '.idf':
        text = path.read_text(encoding='utf-8-sig', errors='replace')
        cleaned = re.sub(r'!.*', '', text)
        counts = Counter(obj.strip().split(',')[0].strip().lower() for obj in cleaned.split(';') if obj.strip())
        if not counts.get('version'):
            raise core.ValidationError('没有找到 IDF Version 对象，请核对文件格式')
        out.update(facts={'objects': dict(counts), 'zones': counts.get('zone', 0)}, next_view='simulation')
        out['warnings'] = ['仅清点对象；未运行 EnergyPlus，不证明可计算。批量导入的 IDF 暂登记为基准候选，多份候选会阻止计算。']
    elif ext in {'.pdf', '.png', '.jpg', '.jpeg'}:
        from .document_reader import worker
        document = worker(path)
        out['warnings'] = document.pop('warnings')
        out['facts'] = document
        out['parser'] = 'local-document-2'
        for page in document['document_pages']:
            locator = f"PDF 第 {page['page']} 页" if page['method'] == 'text' else f"第 {page['page']} 页 · {page['method']}"
            out['candidates'] += text_candidates(page['text'], locator)
            out['candidates'] += text_candidates(page.get('native_text', ''), f"PDF 第 {page['page']} 页文字层")
        out['warnings'].append('OCR / 文本提取仅为待审声明，不代表工程参数准确或已识别空间几何；请逐页对照原图核对。')
        if out['candidates']:
            out['suggested_type'] = 'manual_input'
    elif ext == '.dxf':
        import ezdxf
        drawing = ezdxf.readfile(path)
        out.update(facts={'entities': dict(Counter(e.dxftype() for e in drawing.modelspace())),
                          'units_code': drawing.header.get('$INSUNITS', 0)}, next_view='spatial')
        out['warnings'] = ['CAD 实体已清点，闭合轮廓和图层语义需在建模工作区复核。']
    elif ext in {'.txt', '.csv', '.json'}:
        if path.stat().st_size > 4 * 1024 * 1024:
            raise core.ValidationError('文本资料超过 4 MB，请拆分后导入')
        raw = path.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('gb18030')
        out['facts']['text_preview'] = text[:2400]
        out['candidates'] = text_candidates(text, '文本明确标注字段')
        if ext == '.csv':
            rows = csv.reader(io.StringIO(text))
            out['facts']['columns'] = next(rows, [])[:100]
            out['facts']['data_rows'] = sum(1 for _ in rows)
            out['warnings'].append('已整理表头和行数；账单单位、时段及设备列映射仍需复核。')
        elif ext == '.json':
            obj = json.loads(text)
            if isinstance(obj, dict):
                fields = {**obj, **(obj.get('building') if isinstance(obj.get('building'), dict) else {})}
                labelled = '\n'.join(f'{k}: {v}' for k, v in fields.items() if k in FIELDS and isinstance(v, (str, float, int)) and not isinstance(v, bool))
                out['candidates'] = text_candidates(labelled, 'JSON 顶层 / building 字段')
    elif ext == '.xlsx':
        import zipfile
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                raise core.ValidationError('表格解压大小超过 64 MB，请拆分')
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            out['facts']['sheets'] = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in book.worksheets[:30]]
            out['warnings'] = ['已整理工作表规模；未将任意单元格当成建筑参数。公式仅能读取文件中缓存的值。']
        finally:
            book.close()
    else:
        out['warnings'] = ['原件已保存，本地解析器暂不识别此格式；请转换为可读 PDF、DXF 或 IFC 后继续。']
        out['unsupported'] = True
    return out


class Inbox:
    def __init__(self, domain, auth):
        self.domain, self.store, self.auth = domain, domain.store, auth
        self.lock = threading.RLock()
        self.capacity = threading.BoundedSemaphore(2)
        with self.store.connect() as con:
            con.execute('CREATE TABLE IF NOT EXISTS inbox (id TEXT PRIMARY KEY, owner TEXT NOT NULL, body TEXT NOT NULL)')
            for row in con.execute('SELECT id,body FROM inbox').fetchall():
                body = json.loads(row['body'])
                if body['status'] == 'analyzing':
                    body['status'] = 'interrupted'
                    con.execute('UPDATE inbox SET body=? WHERE id=?', (json.dumps(body, ensure_ascii=False), row['id']))

    def get(self, iid, uid):
        with self.store.connect() as con:
            row = con.execute('SELECT body FROM inbox WHERE id=? AND owner=?', (iid, uid)).fetchone()
        if not row:
            raise HTTPException(404, '资料批次不存在或无访问权限')
        body = json.loads(row[0])
        if body.get('project_id'):
            self.auth.require_project(body['project_id'], uid)
        return body

    def save(self, body):
        with self.store.connect() as con:
            con.execute('UPDATE inbox SET body=? WHERE id=?', (json.dumps(body, ensure_ascii=False, allow_nan=False), body['id']))

    def analyze(self, iid, uid):
        try:
            self._analyze(iid, uid)
        finally:
            self.capacity.release()

    def _analyze(self, iid, uid):
        batch = self.get(iid, uid)
        for source in batch['files']:
            if source.get('corrections'):
                continue  # Preserve the immutable extraction underlying manual revisions.
            source['status'] = 'analyzing'
            self.save(batch)
            try:
                path = core.relative_file(self.store.root, source['path'])
                if core.sha(path) != source['hash']:
                    raise core.ValidationError('原文件校验失败，请重新上传')
                source['analysis'] = parse_source(path)
                source['type'] = source['analysis'].get('suggested_type', source['type'])
                source['status'] = 'unsupported' if source['analysis'].get('unsupported') else 'analyzed'
                source.pop('error', None)
            except Exception as exc:
                source['status'] = 'failed'
                source['error'] = str(exc)[:1000]
            self.save(batch)
        self.summarize(batch)
        self.save(batch)

    def summarize(self, batch):
        candidates = [{**c, 'file_id': s['file_id'], 'file_name': s['name']} for s in batch['files']
                      if s['status'] == 'analyzed' for c in s['analysis']['candidates']]
        suggestions, conflicts = {}, []
        for field in FIELDS:
            values = [c['value'] for c in candidates if c['field'] == field]
            unique = list({json.dumps(v, ensure_ascii=False): v for v in values}.values())
            if len(unique) == 1:
                suggestions[field] = unique[0]
            elif len(unique) > 1:
                conflicts.append(field)
        batch.update(status='review', candidates=candidates, suggestions=suggestions, conflicts=conflicts,
                     missing=[key for key in FIELDS if key not in suggestions])

    def correct(self, iid, uid, body):
        with self.lock:
            batch = self.get(iid, uid)
            if batch['status'] != 'review':
                raise HTTPException(409, '仅可校正尚未归档且已完成分析的批次')
            source = next((s for s in batch['files'] if s['file_id'] == body['file_id']), None)
            if source is None:
                raise HTTPException(404, '文件不存在')
            revision = len(source.get('corrections', []))
            options = body['options']
            if len(options.reason.strip()) < 3:
                raise HTTPException(422, '请填写至少 3 个字符的校正原因')
            if revision != options.expected_revision:
                raise HTTPException(409, '校正记录已变化，请刷新后重试')
            pages = source.get('analysis', {}).get('facts', {}).get('document_pages', [])
            page = next((p for p in pages if p['page'] == options.page), None)
            if page is None:
                raise HTTPException(422, '该页没有解析记录，请重新上传或拆分文件')
            path = core.relative_file(self.store.root, source['path'])
            if core.sha(path) != source['hash']:
                raise HTTPException(409, '原文件校验失败，不能校正旧分析')
            if revision >= 100:
                raise HTTPException(409, '此文件已达到 100 次校正上限，请另建批次保留新资料')
            change = {**options.model_dump(), 'revision': revision + 1, 'actor': uid,
                      'timestamp': core.now(), 'source_hash': source['hash']}
            source.setdefault('corrections', []).append(change)
            # Original page text/boxes remain immutable. Candidates use latest explicit correction.
            latest = {c['page']: c for c in source['corrections']}
            source['analysis']['candidates'] = [c for p in pages for c in text_candidates(
                latest.get(p['page'], p)['text'], f"第 {p['page']} 页 · " +
                (f"人工校正 v{latest[p['page']]['revision']}" if p['page'] in latest else p['method']))]
            for p in pages:
                if p['page'] not in latest:
                    source['analysis']['candidates'] += text_candidates(p.get('native_text', ''), f"PDF 第 {p['page']} 页文字层")
            self.summarize(batch)
            self.save(batch)
            return batch

    def prepare(self, batch, actor):
        """Reuse production tools after atomic adoption; never publish/confirm a model."""
        from .model_schema import ImportOptions
        from .weather import EPWAnalysis
        with self.lock:
            # The caller can repeat commit after a lost response. Re-read persisted
            # progress and reuse an existing model import rather than create duplicates.
            with self.store.connect() as con:
                batch = json.loads(con.execute('SELECT body FROM inbox WHERE id=?', (batch['id'],)).fetchone()[0])
            prepared = batch.setdefault('prepared', {})
            for source in batch['files']:
                fid = source['file_id']
                ext = Path(source['path']).suffix.lower()
                if ext not in {'.epw', '.ifc', '.dxf'} or fid in prepared:
                    continue
                try:
                    if ext == '.epw':
                        result = self.domain.weather.epw(batch['project_id'], EPWAnalysis(file_id=fid))
                        prepared[fid] = {'status': 'analyzed', 'view': 'weather', 'id': result['weather_id']}
                    else:
                        previous = next((r for r in self.store.list('import', batch['project_id']) if r['source']['file_id'] == fid), None)
                        result = previous or self.domain.modeling.start_import(batch['project_id'], ImportOptions(file_id=fid), actor)
                        prepared[fid] = {'status': result['status'], 'view': 'spatial', 'id': result['import_id'], 'error': result.get('error')}
                except Exception as exc:
                    prepared[fid] = {'status': 'failed', 'view': 'weather' if ext == '.epw' else 'spatial', 'error': str(exc)[:1000]}
                self.save(batch)
            return batch

    def commit(self, iid, uid, actor, options):
        with self.lock:
            batch = self.get(iid, uid)
            if batch['status'] == 'committed':
                return batch  # Double click / lost HTTP response cannot duplicate a project.
            if batch['status'] != 'review' or not options.reviewed:
                raise HTTPException(409, '请先完成分析并确认资料来源及建筑声明')
            if not batch.get('project_id') and options.project is None:
                raise HTTPException(422, '新项目必须确认名称、用途、地点、地区和总建筑面积')
            pid = batch.get('project_id') or identifier('project')
            records = []
            if not batch.get('project_id'):
                project = Project(**options.project.model_dump(), project_id=pid, revision=1,
                                  created_at=core.now(), updated_at=core.now()).model_dump(mode='json')
                records.append(('project', pid, project))
            for source in batch['files']:
                path = core.relative_file(self.store.root, source['path'])
                if core.sha(path) != source['hash']:
                    raise HTTPException(409, '资料已变化，请重新导入，不能提交旧分析')
                file = {k: source[k] for k in ('file_id', 'name', 'path', 'hash', 'size')}
                file.update(project_id=pid, revision=1, timestamp=core.now())
                records.append(('file', file['file_id'], file))
                ev = Evidence(evidence_id=identifier('evidence'), project_id=pid, revision=1,
                              timestamp=core.now(), hash=file['hash'], name=file['name'][:160],
                              type=source['type'], source_file=file['file_id'], source_locator='资料收件箱 ' + iid,
                              authority='用户提供，来源真实性待核对', permission='使用许可待复核',
                              acquisition_method='批量上传与本地解析', responsible_person=actor,
                              uncertainty='自动提取结果未经过工程确认',
                              value={'intake_id': iid, 'analysis_status': source['status'], 'analysis': source.get('analysis'),
                                     'corrections': source.get('corrections', []), 'error': source.get('error')},
                              scheme_id='baseline' if source['type'] == 'idf' else None).model_dump(mode='json')
                records.append(('evidence', ev['evidence_id'], ev))
            batch.update(status='committed', project_id=pid, committed_at=core.now())
            with self.store.connect() as con:
                con.execute('BEGIN IMMEDIATE')
                for kind, key, record in records:
                    con.execute('INSERT INTO records VALUES (?,?,?,?,?)', (kind, key, 1, pid, json.dumps(record, ensure_ascii=False, allow_nan=False)))
                con.execute('INSERT OR IGNORE INTO project_owners VALUES (?,?)', (pid, uid))
                con.execute('UPDATE inbox SET body=? WHERE id=?', (json.dumps(batch, ensure_ascii=False, allow_nan=False), iid))
            return batch


def register_intake(app, domain, auth, maximum):
    inbox = app.state.intake = Inbox(domain, auth)

    @app.get('/api/intakes')
    def list_intakes(request: Request):
        with domain.store.connect() as con:
            bodies = [json.loads(r[0]) for r in con.execute('SELECT body FROM inbox WHERE owner=? ORDER BY rowid DESC LIMIT 30', (request.state.user['id'],))]
        return [{'id': b['id'], 'status': b['status'], 'created_at': b['created_at'], 'file_count': len(b['files']), 'project_id': b.get('project_id')} for b in bodies]

    @app.post('/api/intakes', status_code=201)
    def create(body: InboxCreate, request: Request):
        uid = request.state.user['id']
        if body.project_id:
            auth.require_project(body.project_id, uid)
        batch = {'id': identifier('intake'), 'project_id': body.project_id, 'status': 'collecting', 'files': [], 'created_at': core.now()}
        with domain.store.connect() as con:
            if con.execute("SELECT COUNT(*) FROM inbox WHERE owner=? AND json_extract(body,'$.status') != 'committed'", (uid,)).fetchone()[0] >= 30:
                raise HTTPException(409, '最多保留 30 个未提交批次，请完成已有批次')
            con.execute('INSERT INTO inbox VALUES (?,?,?)', (batch['id'], uid, json.dumps(batch)))
        return batch

    @app.get('/api/intakes/{intake_id}')
    def detail(intake_id: str, request: Request):
        return inbox.get(intake_id, request.state.user['id'])

    @app.put('/api/intakes/{intake_id}/draft')
    def draft(intake_id: str, body: InboxDraft, request: Request):
        uid = request.state.user['id']
        with inbox.lock:
            batch = inbox.get(intake_id, uid)
            if batch['status'] != 'review' or batch.get('project_id'):
                raise HTTPException(409, '仅可保存尚未归档的新项目确认字段')
            old = batch.get('draft', {})
            if old.get('revision', 0) != body.expected_revision:
                raise HTTPException(409, '此草稿已在另一页面修改；请重新打开批次核对，当前填写内容仍留在页面中')
            batch['draft'] = {'fields': body.fields, 'revision': body.expected_revision + 1,
                              'saved_at': core.now(), 'actor': uid}
            inbox.save(batch)
            return batch['draft']

    @app.get('/api/intakes/{intake_id}/files/{file_id}/pages/{page}')
    def page_preview(intake_id: str, file_id: str, page: int, request: Request):
        from .document_reader import preview
        batch = inbox.get(intake_id, request.state.user['id'])
        source = next((s for s in batch['files'] if s['file_id'] == file_id), None)
        if not source:
            raise HTTPException(404, '文件不存在')
        path = core.relative_file(domain.store.root, source['path'])
        if path.suffix.lower() not in {'.pdf', '.png', '.jpg', '.jpeg'} or not 1 <= page <= 80:
            raise HTTPException(422, '此文件或页码不支持预览')
        if core.sha(path) != source['hash']:
            raise HTTPException(409, '原文件校验失败')
        return Response(preview(path, page), media_type='image/png', headers={'Cache-Control': 'no-store'})

    @app.post('/api/intakes/{intake_id}/files/{file_id}/corrections')
    def correction(intake_id: str, file_id: str, body: TextCorrection, request: Request):
        return inbox.correct(intake_id, request.state.user['id'], {'file_id': file_id, 'options': body})

    @app.delete('/api/intakes/{intake_id}')
    def discard(intake_id: str, request: Request):
        with inbox.lock:
            batch = inbox.get(intake_id, request.state.user['id'])
            if batch['status'] in {'analyzing', 'committed'}:
                raise HTTPException(409, '分析中的批次或已归档的工程证据不能在这里删除')
            # Only this owner's uncommitted staging artifacts may be removed.
            for source in batch['files']:
                path = core.relative_file(domain.store.root, source['path'])
                if path.parent != domain.store.root / 'inbox-files':
                    raise HTTPException(409, '暂存文件路径异常，未删除')
                path.unlink(missing_ok=True)
            with domain.store.connect() as con:
                con.execute('DELETE FROM inbox WHERE id=? AND owner=?', (intake_id, request.state.user['id']))
        return {'deleted': True}

    @app.post('/api/intakes/{intake_id}/files', status_code=201)
    async def upload(intake_id: str, request: Request, file: UploadFile = File(...)):
        content = await file.read(maximum + 1)
        if not content or len(content) > maximum:
            raise HTTPException(413, '文件为空或超过单文件上传限制')
        extension = Path(file.filename or '').suffix.lower()
        if extension not in TYPES:
            raise HTTPException(422, '不支持此格式；请上传 PDF、IFC、DXF、IDF、EPW、CSV、XLSX、JSON、TXT 或图片')
        with inbox.lock, domain.upload_lock:
            batch = inbox.get(intake_id, request.state.user['id'])
            if batch['status'] not in {'collecting', 'review', 'interrupted'}:
                raise HTTPException(409, '该批次正在分析或已经提交')
            checksum = hashlib.sha256(content).hexdigest()
            if any(s['hash'] == checksum for s in batch['files']):
                return batch
            if len(batch['files']) >= 12 or sum(s['size'] for s in batch['files']) + len(content) > 128 * 1024 * 1024:
                raise HTTPException(413, '每批最多 12 个文件、合计 128 MB')
            fid = identifier('file')
            path = domain.store.root / 'inbox-files' / (fid + extension)
            path.parent.mkdir(exist_ok=True)
            if shutil.disk_usage(path.parent).free < len(content) + 64 * 1024 * 1024:
                raise HTTPException(507, '磁盘空间不足，请由管理员处理')
            try:
                path.write_bytes(content)
                batch['files'].append({'file_id': fid, 'name': Path((file.filename or '').replace('\\', '/')).name[:240],
                                       'path': path.relative_to(domain.store.root).as_posix(), 'size': len(content), 'hash': checksum,
                                       'type': TYPES[extension], 'status': 'uploaded'})
                batch['status'] = 'collecting'
                inbox.save(batch)
            except Exception:
                path.unlink(missing_ok=True)
                raise
        return batch

    @app.post('/api/intakes/{intake_id}/analyze', status_code=202)
    def analyze(intake_id: str, request: Request, background: BackgroundTasks):
        uid = request.state.user['id']
        with inbox.lock:
            batch = inbox.get(intake_id, uid)
            if batch['status'] == 'analyzing':
                return batch
            if batch['status'] == 'committed' or not batch['files']:
                raise HTTPException(409, '请先添加资料；已提交批次不能重新分析')
            if not inbox.capacity.acquire(blocking=False):
                raise HTTPException(409, '分析任务已满，请稍后重试；文件已保存')
            batch['status'] = 'analyzing'
            try:
                inbox.save(batch)
            except Exception:
                inbox.capacity.release()
                raise
        background.add_task(inbox.analyze, intake_id, uid)
        return batch

    @app.post('/api/intakes/{intake_id}/commit')
    def commit(intake_id: str, body: InboxCommit, request: Request):
        user = request.state.user
        batch = inbox.commit(intake_id, user['id'], user['display_name'], body)
        return inbox.prepare(batch, user['display_name'])
